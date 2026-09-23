"""
Fase 1 · Paso 3 — Excel entregable de candidatos.

Genera base_candidatos.xlsx con metadatos, referencia APA7, estatus de
PDF, prioridad (con color por nivel) y columnas para que el investigador
marque inclusión/exclusión. Tres hojas: Candidatos, Escala de prioridad
(desde config) y Resumen.

Lee candidatos_clasificados.json si existe (con prioridad); si no, usa
candidatos_busqueda.json (sin prioridad).

Uso:  python fase1_seleccion/03_generar_excel.py
Salida: data/candidatos/base_candidatos.xlsx
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from litrev_core import config
from litrev_core.estado import registrar_paso
from litrev_core.io import cargar_json
from litrev_core.log import banner, get_logger, timestamp

log = get_logger()

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    log.info("ERROR: openpyxl no instalado. Ejecutar: pip install openpyxl")
    sys.exit(1)


# ── Formato APA7 ─────────────────────────────────────────────────
def _invertir(nombre):
    partes = nombre.strip().split()
    if len(partes) < 2:
        return nombre
    iniciales = " ".join(p[0] + "." for p in partes[:-1] if p)
    return f"{partes[-1]}, {iniciales}"


def autores_apa(autores):
    if not autores:
        return "Sin autor"
    if len(autores) == 1:
        return _invertir(autores[0])
    if len(autores) == 2:
        return f"{_invertir(autores[0])} & {_invertir(autores[1])}"
    if len(autores) <= 20:
        return ", ".join(_invertir(a) for a in autores[:-1]) + ", & " + _invertir(autores[-1])
    return ", ".join(_invertir(a) for a in autores[:19]) + ", ... " + _invertir(autores[-1])


def apa7(art):
    ref = f"{autores_apa(art.get('autores', []))} ({art.get('anio', 's.f.')}). "
    titulo = art.get("titulo", "Sin título") or "Sin título"
    ref += titulo if titulo.endswith(".") else titulo + "."
    if art.get("revista"):
        ref += f" {art['revista']}."
    doi = art.get("doi", "")
    if doi:
        ref += " " + (doi if doi.startswith("http") else f"https://doi.org/{doi}")
    return ref


def primer_autor(autores):
    if not autores:
        return "Sin autor"
    partes = autores[0].strip().split()
    return partes[-1] if partes else autores[0]


def main():
    banner(log, "LITREV-AI · Fase 1 · Excel de candidatos")

    paths = config.get_paths()
    fuente = paths.candidatos_clasificados if paths.candidatos_clasificados.exists() \
        else paths.candidatos_busqueda
    if not fuente.exists():
        log.info("No hay candidatos. Ejecutar primero 01_buscar.py (y 02_clasificar_prioridad.py).")
        sys.exit(1)

    articulos = cargar_json(fuente)
    pdfs = {f.stem for f in paths.pdf_dir.glob("*.pdf")} if paths.pdf_dir.exists() else set()
    niveles = config.cfg_prioridad().get("niveles", {})
    proy = config.cfg_proyecto().get("proyecto", {})

    log.info(f"Candidatos: {len(articulos)}  | PDFs descargados: {len(pdfs)}  | fuente: {fuente.name}")

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=True)
    borde = Border(*[Side(style="thin")] * 4)
    fills_nivel = {n: PatternFill("solid", fgColor=(niveles[n] or {}).get("color", "FFFFFF"))
                   for n in niveles}

    wb = Workbook()

    # ── Hoja 1: Candidatos ──
    ws = wb.active
    ws.title = "Candidatos"
    columnas = [
        ("ID", 18), ("Primer autor", 16), ("Autores (todos)", 38), ("Año", 7),
        ("Título", 50), ("Revista", 28), ("DOI", 28), ("Referencia APA7", 60),
        ("Abstract", 70), ("País", 16), ("Idioma", 8), ("Tipo", 14),
        ("Citaciones", 10), ("Open Access", 11), ("Estatus PDF", 13),
        ("Fuente API", 12), ("Query origen", 11), ("Prioridad", 10),
        ("Score", 8), ("Confianza prioridad", 14), ("Razón prioridad", 40),
        ("Decisión inclusión", 16), ("Notas investigadora", 36),
    ]
    for c, (nombre, ancho) in enumerate(columnas, 1):
        cell = ws.cell(1, c, nombre)
        cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, header_align, borde
        ws.column_dimensions[get_column_letter(c)].width = ancho

    col_prioridad = 18
    for r, art in enumerate(articulos, 2):
        aid = art.get("id_candidato", "")
        estatus = "Descargado" if aid in pdfs else ("Sin URL" if not art.get("pdf_url") else "Pendiente")
        valores = [
            aid, primer_autor(art.get("autores", [])),
            "; ".join(art.get("autores", []) or []) or "Sin autor",
            art.get("anio", "no_reportado"), art.get("titulo", "") or "no_reportado",
            art.get("revista", "") or "no_reportado", art.get("doi", "") or "no_reportado",
            apa7(art), art.get("abstract", "") or "no_reportado",
            art.get("pais", "") or "no_reportado", art.get("idioma", "") or "no_reportado",
            art.get("tipo", "") or "no_reportado", art.get("citaciones", 0),
            "Sí" if art.get("is_oa") else "No", estatus,
            art.get("fuente_api", ""), art.get("query_origen", ""),
            art.get("prioridad", ""), art.get("prioridad_score", ""),
            art.get("prioridad_confianza", ""), art.get("prioridad_razon", ""),
            "", "",
        ]
        for c, valor in enumerate(valores, 1):
            cell = ws.cell(r, c, valor)
            cell.alignment, cell.border = cell_align, borde
        pri = art.get("prioridad", "")
        if pri in fills_nivel:
            ws.cell(r, col_prioridad).fill = fills_nivel[pri]
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(articulos) + 1}"
    ws.freeze_panes = "A2"

    # ── Hoja 2: Escala de prioridad (desde config) ──
    ws2 = wb.create_sheet("Escala de prioridad")
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 80
    ws2.column_dimensions["C"].width = 48
    ws2.cell(1, 1, "Escala de prioridad — LITREV-AI").font = Font(bold=True, size=12, color="2F5496")
    ws2.cell(2, 1, f"Proyecto: {proy.get('nombre', '')}").font = Font(size=10)
    ws2.cell(3, 1, f"Generado: {timestamp()}").font = Font(size=10)
    for c, nombre in enumerate(["Nivel", "Descripción", "Acción recomendada"], 1):
        cell = ws2.cell(5, c, nombre)
        cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, header_align, borde
    for i, (nivel, datos) in enumerate(niveles.items()):
        row = 6 + i
        datos = datos or {}
        etiqueta = f"{nivel} — {datos.get('label', '')}"
        for c, valor in enumerate([etiqueta, datos.get("descripcion", ""), datos.get("accion", "")], 1):
            cell = ws2.cell(row, c, valor)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = borde
            cell.fill = fills_nivel.get(nivel, PatternFill())
        ws2.row_dimensions[row].height = 70

    # ── Hoja 3: Resumen ──
    ws3 = wb.create_sheet("Resumen")
    ws3.column_dimensions["A"].width = 32
    ws3.column_dimensions["B"].width = 14
    ws3.cell(1, 1, "Resumen del corpus de candidatos").font = Font(bold=True, size=12, color="2F5496")
    con_pdf = sum(1 for a in articulos if a.get("id_candidato") in pdfs)
    stats = [
        ("Total candidatos", len(articulos)),
        ("Con PDF descargado", con_pdf),
        ("Sin URL de PDF", sum(1 for a in articulos if not a.get("pdf_url"))),
        ("", ""),
        ("Fuente: OpenAlex", sum(1 for a in articulos if a.get("fuente_api") == "openalex")),
        ("Fuente: Semantic Scholar", sum(1 for a in articulos if a.get("fuente_api") == "semantic_scholar")),
        ("Open Access", sum(1 for a in articulos if a.get("is_oa"))),
        ("", ""),
    ]
    for nivel in niveles:
        n = sum(1 for a in articulos if a.get("prioridad") == nivel)
        stats.append((f"Prioridad {nivel} — {(niveles[nivel] or {}).get('label', '')}", n))
    stats.append(("Dudosos (confianza)", sum(1 for a in articulos if a.get("prioridad_confianza") == "dudosa")))
    for i, (label, valor) in enumerate(stats, 3):
        ws3.cell(i, 1, label).font = Font(size=10)
        if valor != "":
            ws3.cell(i, 2, valor).font = Font(size=10)

    wb.save(paths.base_candidatos_xlsx)
    log.info(f"\nExcel generado: {paths.base_candidatos_xlsx}")
    log.info("Hojas: Candidatos · Escala de prioridad · Resumen")

    registrar_paso(paths.estado_json, "fase1", "excel",
                   {"candidatos": len(articulos), "con_pdf": con_pdf})


if __name__ == "__main__":
    main()
