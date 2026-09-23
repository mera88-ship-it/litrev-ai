"""
Fase 1 · Paso 5 — Cierre de selección (aprobación).

Lee la columna "Decisión inclusión" del Excel revisado por el investigador
y construye candidatos_aprobados.json: el contrato de entrada de la Fase 2.
Se aceptan como inclusión: incluir / sí / si / include / x / 1.

La CONVERSIÓN a Markdown ya NO ocurre aquí: es el primer paso de la
Fase 2 (procesamiento), porque convertir a un formato trabajable es el
inicio del procesamiento de la literatura, no de su selección.

Uso:
  python fase1_seleccion/05_aprobar.py
  python fase1_seleccion/05_aprobar.py --con-pdf   # aprueba los que tienen PDF (sin Excel)
Salida: data/candidatos/candidatos_aprobados.json
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from litrev_core import config
from litrev_core.estado import registrar_paso
from litrev_core.io import cargar_json, guardar_json
from litrev_core.log import banner, get_logger

log = get_logger()

INCLUIR = {"incluir", "sí", "si", "include", "x", "1", "true", "verdadero"}


def leer_decisiones(xlsx_path):
    """Devuelve {id_candidato: decision_en_minuscula} desde el Excel."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        log.info("openpyxl no instalado; no se pueden leer decisiones del Excel.")
        return {}
    if not xlsx_path.exists():
        return {}
    wb = load_workbook(xlsx_path, read_only=True)
    ws = wb["Candidatos"] if "Candidatos" in wb.sheetnames else wb.active
    cabecera = next(ws.iter_rows(min_row=1, max_row=1))
    col_id = col_dec = None
    for c, cell in enumerate(cabecera, 1):
        if cell.value == "ID":
            col_id = c
        elif cell.value == "Decisión inclusión":
            col_dec = c
    decisiones = {}
    if col_id and col_dec:
        for row in ws.iter_rows(min_row=2):
            aid = row[col_id - 1].value
            dec = row[col_dec - 1].value or ""
            if aid:
                decisiones[str(aid)] = str(dec).strip().lower()
    wb.close()
    return decisiones


def main():
    banner(log, "LITREV-AI · Fase 1 · Aprobación de candidatos")
    con_pdf_flag = "--con-pdf" in sys.argv[1:]

    paths = config.get_paths()
    fuente = paths.candidatos_clasificados if paths.candidatos_clasificados.exists() \
        else paths.candidatos_busqueda
    if not fuente.exists():
        log.info("No hay candidatos. Ejecutar los pasos anteriores de la Fase 1.")
        sys.exit(1)

    articulos = cargar_json(fuente)
    por_id = {a.get("id_candidato"): a for a in articulos}
    pdfs = {f.stem for f in paths.pdf_dir.glob("*.pdf")} if paths.pdf_dir.exists() else set()

    decisiones = leer_decisiones(paths.base_candidatos_xlsx)
    marcados = {aid for aid, d in decisiones.items() if d}

    if con_pdf_flag:
        aprobados_ids = [aid for aid in por_id if aid in pdfs]
        log.info(f"Modo --con-pdf: {len(aprobados_ids)} candidatos con PDF descargado.")
    elif marcados:
        aprobados_ids = [aid for aid, d in decisiones.items() if d in INCLUIR and aid in por_id]
        log.info(f"Decisiones leídas: {len(marcados)} marcadas → {len(aprobados_ids)} incluidas.")
    else:
        log.info("⚠ El Excel no tiene decisiones en 'Decisión inclusión'.")
        log.info("  Marcar la inclusión en el Excel y reejecutar, o usar --con-pdf como recurso.")
        sys.exit(1)

    aprobados = [por_id[aid] for aid in aprobados_ids if aid in por_id]
    con = sum(1 for a in aprobados if a.get("id_candidato") in pdfs)
    guardar_json(paths.candidatos_aprobados, aprobados, backup=True)

    log.info("\n" + "=" * 64)
    log.info(f"Aprobados: {len(aprobados)}  (con PDF descargado: {con} | sin PDF aún: {len(aprobados) - con})")
    log.info(f"Contrato para la Fase 2: {paths.candidatos_aprobados}")
    log.info("Siguiente: Fase 2 · conversión a Markdown (fase2_procesamiento/01_convertir_markdown.py)")
    log.info("=" * 64)

    registrar_paso(paths.estado_json, "fase1", "aprobar",
                   {"aprobados": len(aprobados), "con_pdf": con})


if __name__ == "__main__":
    main()
